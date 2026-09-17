"""Edge cases for the Markdown, PDF, and file-reading tools and the WeasyPrint converter."""
from pathlib import Path
from types import SimpleNamespace

import pandas as pd
import pytest
from docx import Document

from api.context import reset_session_context, set_session_context
from tools import markdown_tools, pdf_tools, upload_file_read_tool
from tools.markdown_tools import generate_markdown
from tools.pdf_tools import convert_md_to_pdf
from tools.upload_file_read_tool import read_file_content
from utils import word_converter


@pytest.fixture
def session(tmp_path):
    token = set_session_context(str(tmp_path.resolve()))
    yield tmp_path.resolve()
    reset_session_context(token)


# --------------------------------------------------------- generate_markdown --
def test_generate_markdown_keeps_existing_suffix_and_reports(session, monitor_calls):
    result = generate_markdown.invoke({"content": "# Title", "filename": "report.md"})
    assert result == f"Markdown file '{session / 'report.md'}' generated and saved successfully."
    assert (session / "report.md").read_text(encoding="utf-8") == "# Title"
    assert not (session / "report.md.md").exists()
    assert monitor_calls == [("Markdown generation tool", {"content": "# Title"})]


@pytest.mark.parametrize("path", ["", "."])
def test_generate_markdown_treats_dot_and_empty_path_as_session_root(session, path):
    generate_markdown.invoke({"content": "x", "filename": f"root_{len(path)}", "path": path})
    assert (session / f"root_{len(path)}.md").read_text(encoding="utf-8") == "x"


def test_generate_markdown_accepts_absolute_path_inside_session(session):
    generate_markdown.invoke({"content": "abs", "filename": "abs", "path": str(session / "deep" / "er")})
    assert (session / "deep" / "er" / "abs.md").read_text(encoding="utf-8") == "abs"


def test_generate_markdown_refuses_to_escape_the_session(session, tmp_path):
    with pytest.raises(ValueError, match="Only this conversation's non-hidden files"):
        generate_markdown.invoke({"content": "leak", "filename": "../escape"})
    assert not (session.parent / "escape.md").exists()
    with pytest.raises(ValueError):
        generate_markdown.invoke({"content": "leak", "filename": ".state.json"})


def test_generate_markdown_reports_write_failures(session):
    (session / "blocker").write_text("I am a file, not a directory", encoding="utf-8")
    result = generate_markdown.invoke({"content": "x", "filename": "report", "path": "blocker"})
    assert result.startswith("Failed to generate Markdown file:")


# --------------------------------------------------------- convert_md_to_pdf --
def test_convert_md_to_pdf_defaults_output_next_to_source(session, monitor_calls):
    (session / "report.md").write_text("# Report\n\n| a | b |\n|---|---|\n| 1 | 2 |\n", encoding="utf-8")
    result = convert_md_to_pdf.invoke({"md_filename": "report"})
    assert result.startswith("Converted successfully")
    assert (session / "report.pdf").read_bytes().startswith(b"%PDF-")
    assert monitor_calls == [("Markdown-to-PDF tool", None)]


def test_convert_md_to_pdf_adds_pdf_suffix_to_requested_name(session):
    (session / "report.md").write_text("# Report", encoding="utf-8")
    convert_md_to_pdf.invoke({"md_filename": "report.md", "pdf_filename": "final"})
    assert (session / "final.pdf").exists()


def test_convert_md_to_pdf_reports_escaping_paths_as_failures(session, tmp_path):
    (tmp_path.parent / "outside.md").write_text("# Outside", encoding="utf-8")
    result = convert_md_to_pdf.invoke({"md_filename": "../outside.md"})
    assert result.startswith("Conversion failed:") and "Only this conversation" in result
    (session / "report.md").write_text("# Report", encoding="utf-8")
    result = convert_md_to_pdf.invoke({"md_filename": "report.md", "pdf_filename": "../escape.pdf"})
    assert result.startswith("Conversion failed:")
    assert not (tmp_path.parent / "escape.pdf").exists()


def test_convert_md_to_pdf_wraps_renderer_errors(session, monkeypatch):
    (session / "report.md").write_text("# Report", encoding="utf-8")

    def broken(*args, **kwargs):
        raise RuntimeError("renderer exploded")

    monkeypatch.setattr(word_converter, "convert_md_to_pdf", broken)
    result = convert_md_to_pdf.invoke({"md_filename": "report.md"})
    assert result == "Conversion failed: renderer exploded"


# ------------------------------------------------------------ word_converter --
def test_word_converter_renders_tables_and_code_fences(tmp_path, monkeypatch):
    captured = {}

    class FakeHTML:
        def __init__(self, string):
            captured["html"] = string

        def write_pdf(self, target):
            Path(target).write_bytes(b"%PDF-fake")

    monkeypatch.setattr(word_converter, "HTML", FakeHTML)
    source = tmp_path / "doc.md"
    source.write_text("# H\n\n| a |\n|---|\n| 1 |\n\n```python\nprint(1)\n```\n", encoding="utf-8")
    result = word_converter.convert_md_to_pdf(source, tmp_path / "doc.pdf")
    assert result.startswith("Converted successfully") and "WeasyPrint" in result
    assert "<table>" in captured["html"] and "<code" in captured["html"]
    assert 'charset="UTF-8"' in captured["html"]


def test_word_converter_reports_missing_output(tmp_path, monkeypatch):
    class SilentHTML:
        def __init__(self, string):
            pass

        def write_pdf(self, target):
            return None

    monkeypatch.setattr(word_converter, "HTML", SilentHTML)
    source = tmp_path / "doc.md"
    source.write_text("# H", encoding="utf-8")
    result = word_converter.convert_md_to_pdf(source, tmp_path / "doc.pdf")
    assert result.startswith("Conversion finished but file was not created")


def test_word_converter_reports_read_failures(tmp_path):
    result = word_converter.convert_md_to_pdf(tmp_path / "missing.md", tmp_path / "missing.pdf")
    assert result.startswith("Conversion failed:")


# ---------------------------------------------------------- read_file_content --
def test_read_file_content_reads_text_variants(session, monitor_calls):
    (session / "notes.txt").write_text("plain", encoding="utf-8")
    (session / "UPPER.MD").write_text("# upper", encoding="utf-8")
    (session / "data.csv").write_text("a,b\n1,2", encoding="utf-8")
    assert read_file_content.invoke({"filename": "notes.txt", "instruction": "summarize"}) == "plain"
    assert read_file_content.invoke({"filename": "UPPER.MD"}) == "# upper"
    assert read_file_content.invoke({"filename": "data.csv"}) == "a,b\n1,2"
    assert monitor_calls[0] == ("File content reader", {"filename": "notes.txt", "instruction": "summarize"})
    assert monitor_calls[1][1]["instruction"] == "Extract all content"


def test_read_file_content_rejects_binary_unknown_formats(session):
    (session / "blob.dat").write_bytes(b"\xff\xfe\x00binary")
    assert read_file_content.invoke({"filename": "blob.dat"}) == (
        "Error: unsupported file format '.dat'; cannot read it as text.")


def test_read_file_content_refuses_to_escape_the_session(session):
    with pytest.raises(ValueError, match="Only this conversation's non-hidden files"):
        read_file_content.invoke({"filename": "../secret.md"})
    with pytest.raises(ValueError):
        read_file_content.invoke({"filename": ".state.json"})


@pytest.mark.parametrize("name,attribute,message", [
    ("doc.docx", "docx", "install 'python-docx'"),
    ("doc.pdf", "pypdf", "install 'pypdf'"),
    ("doc.xlsx", "pd", "install 'pandas'"),
])
def test_read_file_content_reports_missing_optional_readers(session, monkeypatch, name, attribute, message):
    (session / name).write_bytes(b"placeholder")
    monkeypatch.setattr(upload_file_read_tool, attribute, None)
    assert message in read_file_content.invoke({"filename": name})


@pytest.mark.parametrize("name,prefix", [
    ("bad.docx", "Failed to read file:"),
    ("bad.pdf", "Failed to read file:"),
    ("bad.xlsx", "Failed to read Excel:"),
])
def test_read_file_content_reports_corrupt_documents(session, name, prefix):
    (session / name).write_bytes(b"garbage")
    assert read_file_content.invoke({"filename": name}).startswith(prefix)


def test_read_file_content_excel_summary(session):
    pd.DataFrame({"region": ["East", "West", "North"], "sales": [42, 18, 7]}).to_excel(
        session / "sales.xlsx", index=False)
    out = read_file_content.invoke({"filename": "sales.xlsx"})
    assert "File: sales.xlsx" in out
    assert "Rows: 3, Columns: 2" in out
    assert "Column names: region, sales" in out
    assert "[First 5 rows]:" in out and "North" in out
    assert "[Descriptive statistics]:" in out and "mean" in out


def test_read_file_content_word_paragraphs_joined_with_newlines(session):
    document = Document()
    document.add_paragraph("First paragraph")
    document.add_paragraph("Second paragraph")
    document.save(session / "doc.docx")
    assert read_file_content.invoke({"filename": "doc.docx"}) == "First paragraph\nSecond paragraph"


def test_read_file_content_pdf_pages_without_text_are_tolerated(session, monkeypatch):
    (session / "scan.pdf").write_bytes(b"placeholder")
    pages = [SimpleNamespace(extract_text=lambda: None), SimpleNamespace(extract_text=lambda: "Page two")]
    monkeypatch.setattr(upload_file_read_tool, "pypdf",
                        SimpleNamespace(PdfReader=lambda path: SimpleNamespace(pages=pages)))
    assert read_file_content.invoke({"filename": "scan.pdf"}) == "\nPage two"
